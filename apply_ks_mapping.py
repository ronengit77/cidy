"""
Rewrite Jack/Jon knowledge-source IDs in a test yaml to their production
equivalents, using KS_Mapping_Table.xlsx, before pasting into Copilot Studio.

Usage:

    py apply_ks_mapping.py "Jack_PD_Tests\\Formulate_Response_PD_Broad_to_Narrow.yaml"

Writes "<same folder>\\<stem>.production_ready.yaml" with each knowledge-source
ID replaced by its mapped production ID (only for rows in the mapping table
where ProductionKnowledgeSourceID is filled in). The original input file is
never modified. Any knowledge-source reference with no confirmed mapping is
left untouched in the output and printed as a warning - never guessed.
"""

import argparse
import os
import re
import sys

try:
    import openpyxl
except ImportError as e:
    print(f"Missing dependency: {e}")
    print("Install requirements with:\n  py -m pip install openpyxl")
    sys.exit(1)

ROOT = r"C:\Users\RRAPOPOR\Documents\Cidy testing"
DEFAULT_MAPPING_XLSX = os.path.join(ROOT, "KS_Mapping_Table.xlsx")
KS_LINE_RE = re.compile(r"^(\s*-\s+)(copilots_header_3141e\.topic\.\S+)(\s*)$")


def load_mapping(xlsx_path):
    if not os.path.exists(xlsx_path):
        print(f"STOP: mapping table not found: {xlsx_path}")
        print("Run build_ks_mapping_table.py first.")
        sys.exit(1)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "Mapping" not in wb.sheetnames:
        print(f"STOP: 'Mapping' sheet not found in {xlsx_path} (found {wb.sheetnames})")
        sys.exit(1)
    ws = wb["Mapping"]
    headers = [c.value for c in ws[1]]
    required = ["TestKnowledgeSourceID", "ProductionKnowledgeSourceID"]
    missing = [h for h in required if h not in headers]
    if missing:
        print(f"STOP: mapping table missing expected columns: {missing}")
        sys.exit(1)

    mapping = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        test_id = d.get("TestKnowledgeSourceID")
        prod_id = d.get("ProductionKnowledgeSourceID")
        if test_id and prod_id:
            mapping[test_id.strip()] = prod_id.strip()
    return mapping


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_yaml", help="Path to a Jack/Jon test yaml file")
    parser.add_argument("--mapping", default=DEFAULT_MAPPING_XLSX, help="Path to KS_Mapping_Table.xlsx")
    args = parser.parse_args()

    in_path = args.input_yaml
    if not os.path.exists(in_path):
        print(f"STOP: input file not found: {in_path}")
        sys.exit(1)

    mapping = load_mapping(args.mapping)
    print(f"Loaded {len(mapping)} confirmed mapping(s) from {os.path.basename(args.mapping)}")

    with open(in_path, encoding="utf-8", newline="") as f:
        lines = f.readlines()

    out_lines = []
    replaced = 0
    unmapped = []
    for i, line in enumerate(lines):
        m = KS_LINE_RE.match(line.rstrip("\n"))
        if m:
            lead, test_id, trail = m.groups()
            if test_id in mapping:
                prod_id = mapping[test_id]
                out_lines.append(f"{lead}{prod_id}{trail}\n")
                replaced += 1
                continue
            else:
                unmapped.append((i + 1, test_id))
        out_lines.append(line)

    stem, ext = os.path.splitext(in_path)
    out_path = f"{stem}.production_ready{ext}"
    with open(out_path, "w", encoding="utf-8", newline="") as f:
        f.writelines(out_lines)

    print(f"\nInput : {in_path}")
    print(f"Output: {out_path}  (original input file left unchanged)")
    print(f"\nKnowledge-source references replaced: {replaced}")
    if unmapped:
        print(f"\nWARNING: {len(unmapped)} reference(s) left UNCHANGED (no confirmed mapping found):")
        for line_no, test_id in unmapped:
            print(f"  line {line_no}: {test_id}")
        print("\n-> these still need a manual knowledge-source reselection in Copilot Studio,")
        print("   or fill them into KS_Mapping_Table.xlsx and re-run this script.")
    else:
        print("\nNo unmapped references - every knowledge-source ID in this file had a confirmed mapping.")


if __name__ == "__main__":
    main()
