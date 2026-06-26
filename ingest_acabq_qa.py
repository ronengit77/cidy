"""
Ingest new ACABQ Q&A docx/pdf files into ACABQ_QA_Complete_Tagged.xlsx.

Normal usage (run any time after dropping new files into the Ingest folder):

    py ingest_acabq_qa.py

Workflow:
  1. Scans Knowledge/Programme Development/ACABQ/Ingest for .docx/.pdf files.
  2. For each file, auto-detects which of the 3 known ACABQ Q&A structures it
     uses (docx direct "N." marker, docx/pdf "Question N." marker, or old-style
     PDF memo with numbered sub-answers) and extracts Question/Answer pairs.
     If no structure can be confidently detected, the file is left in Ingest
     and flagged for manual review - it is NEVER force-parsed or guessed at.
  3. Drops any (question, answer) pair that's an exact duplicate of a row
     already in the workbook (or already added earlier in the same run).
  4. Auto-tags Category (9-bucket keyword heuristic) and isDESA ("DESA"
     literal match in the Answer) the same way the rest of the workbook was
     tagged. These are heuristic, not manually reviewed - spot-check them.
  5. Year is parsed from the filename (first 19xx/20xx 4-digit run found),
     else left blank - never guessed from today's date. FilePathFull/
     FolderPath point at the local Ingested-folder location (file:// link);
     there is no SharePoint URL or DocumentID available for brand-new files,
     so those are left blank rather than invented.
  6. Appends new rows to the QA table (extending the Table object and the
     Category/isDESA dropdown validation ranges so the workbook stays a
     genuine Excel Table), adds a SourceFile hyperlink per new row, and logs
     a dated entry in the Log sheet describing exactly what happened.
  7. Moves each file that was successfully parsed (even if it produced zero
     *new* rows because everything in it was already known) to the Ingested
     folder. Files that were skipped (duplicate filename / unparseable /
     scanned PDF with no text layer) are left in Ingest untouched.

Never alters existing rows, never invents Question/Answer text, never
force-parses a file it isn't confident about.
"""

import argparse
import datetime
import os
import re
import shutil
import sys
from pathlib import Path

try:
    import docx
    import pdfplumber
    import openpyxl
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as e:
    print(f"Missing dependency: {e}")
    print("Install requirements with:\n  py -m pip install python-docx pdfplumber openpyxl")
    sys.exit(1)

ACABQ_DIR = r"C:\Users\RRAPOPOR\Documents\Cidy testing\Knowledge\Programme Development\ACABQ"
DEFAULT_INGEST_DIR = os.path.join(ACABQ_DIR, "Ingest")
DEFAULT_INGESTED_DIR = os.path.join(ACABQ_DIR, "Ingested")
DEFAULT_XLSX = r"C:\Users\RRAPOPOR\Documents\Cidy testing\ACABQ_QA_Complete_Tagged.xlsx"
SHEET_NAME = "QA"
TABLE_NAME = "ACABQ_QA_Tagged"
DEFAULT_MIN_MARKERS = 2  # minimum marker hits required to trust an auto-detected structure

CATEGORIES = [
    "Budget & Expenditure",
    "Staffing & Consultants",
    "Programme Strategy & Governance",
    "Requests, Country Engagement & Outreach",
    "Monitoring, Evaluation & Reporting",
    "Projects & Field Activities",
    "COVID-19 Impact",
    "Human Rights & Gender",
    "Other / Administrative",
]

CATEGORY_RULES = [
    ("COVID-19 Impact", r"covid"),
    ("Human Rights & Gender", r"human right|gender|\bwomen\b|indigenous|universal periodic review|\bupr\b|albinism|death penalty|\bnhri"),
    ("Staffing & Consultants", r"\bgta\b|general temporary assistance|consultant|organigram|org chart|chief of (service|unit)|inter-regional advis|\bp-4\b|\bp-5\b|\bd-1\b|gs \(ol\)|recruitment|incumbent|vacant position|staff cost|construction and maintenance worker|\bdriver\b|mail assistant|finance assistant|procurement assistant|budget assistant|information system(s)? assistant|team assistant|\bposts?\b"),
    ("Projects & Field Activities", r"field project|field activit|list of projects|project execut|pilot project"),
    ("Requests, Country Engagement & Outreach", r"request(s)? for (assistance|support)|resident coordinator|\brc\b|\brco\b|outreach|demand[- ]driven|awareness|applications are made"),
    ("Monitoring, Evaluation & Reporting", r"progress report|evaluation|\boios\b|performance measure|monitoring|assessment|lessons learned|common reporting standard|\bimpact\b"),
    ("Programme Strategy & Governance", r"development account|legislative|mandate|resolution \d|policy decision|decision process|allocation of resources|governance|\bstrategy\b|acabq (recommend|comment)|quadrennial|\bqcpr\b"),
    ("Budget & Expenditure", r"expenditure|budget|recost|appropriat|variance|table 23|table s\.23|\bcost\b|resource requirement|fascicle"),
]


def classify_category(q, a):
    text = ((q or "") + " " + (a or "")).lower()
    for label, pattern in CATEGORY_RULES:
        if re.search(pattern, text):
            return label
    return "Other / Administrative"


def classify_is_desa(a):
    return "Yes" if "DESA" in (a or "") else "No"


def extract_year_from_filename(fname):
    m = re.search(r"(19|20)\d{2}", fname)
    return m.group(0) if m else ""


def local_file_uri(abs_path):
    try:
        return Path(abs_path).resolve().as_uri()
    except Exception:
        return ""


def normalize_text(s):
    return " ".join(str(s).split()).strip().lower()


def norm_key(q, a):
    return (normalize_text(q), normalize_text(a))


# ---------------------------------------------------------------------------
# Extraction helpers (same logic validated on the original 19-file batch)
# ---------------------------------------------------------------------------

def extract_docx_paragraphs(path):
    d = docx.Document(path)
    paras = []
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            from docx.text.paragraph import Paragraph
            para = Paragraph(child, d)
            paras.append(para.text)
        elif tag == "tbl":
            from docx.table import Table as DocxTable
            tbl = DocxTable(child, d)
            for row in tbl.rows:
                cells = [c.text for c in row.cells]
                paras.append(" | ".join(cells))
    return paras


PAGE_BOILERPLATE_LINE_RES = [
    re.compile(r"^\d{1,3}$"),
    re.compile(r"^UNITED NATIONS.*INTEROFFICE MEMORANDUM.*MEMORANDUM INTERIEUR.*PAGE\s*\d*$", re.IGNORECASE),
    re.compile(r"^UNAMID$", re.IGNORECASE),
    re.compile(r"^Response No\.?\s*\d+$", re.IGNORECASE),
    re.compile(r"^Section\s+\d+[A-Za-z]*$", re.IGNORECASE),
    re.compile(r"^Set\s+\d+(\s+of\s+\d+)?$", re.IGNORECASE),
]


def strip_page_boilerplate(pages):
    cleaned_pages = []
    for i, page_text in enumerate(pages):
        if i == 0:
            cleaned_pages.append(page_text)
            continue
        lines = page_text.split("\n")
        idx = 0
        while idx < len(lines):
            line = lines[idx].strip()
            if line == "" or any(r.match(line) for r in PAGE_BOILERPLATE_LINE_RES):
                idx += 1
                continue
            break
        cleaned_pages.append("\n".join(lines[idx:]))
    return cleaned_pages


def extract_pdf_text(path):
    raw_pages = []
    has_any_text = False
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            if t.strip():
                has_any_text = True
            raw_pages.append(t)
    cleaned_pages = strip_page_boilerplate(raw_pages)
    return "\n".join(cleaned_pages), has_any_text


MARKER_RE_A = re.compile(r"^(\d{1,3}[a-z]?)\.\s+")
MARKER_RE_B = re.compile(r"^Question\s+(\d{1,3})\.\s*")
TOP_RE_C = re.compile(r"(?m)^(\d{1,3}[a-z]?)\.\s+(?=[A-Z(])")
SUB_RE_C = re.compile(r"(?m)^(\d{1,3}[a-z]?)\.(\d{1,2})\s+")
CC_RE_C = re.compile(r"(?m)^cc:|secretariat of the fifth committee", re.IGNORECASE)


def parse_marker_paragraphs(paragraphs, marker_re):
    """Generic 'marker at start of paragraph -> answer is everything until the
    next marker' parser. Used for both family A (bare 'N.') and family B
    ('Question N.') since the surrounding logic is identical."""
    marker_indices = [i for i, p in enumerate(paragraphs) if marker_re.match(p.strip())]
    warnings = []
    pairs = []
    if not marker_indices:
        return pairs, warnings, 0

    for idx, start in enumerate(marker_indices):
        end = marker_indices[idx + 1] if idx + 1 < len(marker_indices) else len(paragraphs)
        question_raw = paragraphs[start].strip()
        question = marker_re.sub("", question_raw, count=1).strip()

        answer_paras = paragraphs[start + 1:end]
        while answer_paras and not answer_paras[0].strip():
            answer_paras.pop(0)
        while answer_paras and not answer_paras[-1].strip():
            answer_paras.pop()
        answer = "\n".join(answer_paras).strip()

        if not question:
            warnings.append(f"empty_question_at_index_{start}")
            continue
        if not answer:
            warnings.append(f"empty_answer_at_index_{start}")

        pairs.append((question, answer))

    return pairs, warnings, len(marker_indices)


def find_question_answer_split(collapsed_text):
    pos = 0
    boundary = None
    while True:
        qmark = collapsed_text.find("?", pos)
        if qmark == -1:
            break
        boundary = qmark + 1
        rest = collapsed_text[boundary:].lstrip()
        next_q = rest.find("?")
        next_period = rest.find(". ")
        next_end = next_q if next_q != -1 else -1
        if next_end != -1 and (next_period == -1 or next_end < next_period) and next_end < 80:
            pos = boundary
            continue
        break

    if boundary is not None:
        return collapsed_text[:boundary], collapsed_text[boundary:]

    pos = 0
    while True:
        m = re.search(r"\.\s+", collapsed_text[pos:])
        if not m:
            break
        cand = pos + m.end()
        rest = collapsed_text[cand:]
        m2 = re.search(r"\.\s+", rest)
        seg_len = m2.start() if m2 else len(rest)
        if seg_len < 80 and rest[:1].isupper():
            pos = cand
            continue
        boundary = cand
        break

    if boundary is not None:
        return collapsed_text[:boundary], collapsed_text[boundary:]

    return collapsed_text, ""


def parse_family_c(text):
    warnings = []
    pairs = []

    cutoff = 0
    cc_match = CC_RE_C.search(text)
    if cc_match and cc_match.start() < 3000:
        cutoff = cc_match.end()

    body = text[cutoff:]

    top_matches = list(TOP_RE_C.finditer(body))
    if not top_matches:
        return pairs, warnings

    sub_matches = list(SUB_RE_C.finditer(body))
    subs_by_id = {}
    for m in sub_matches:
        subs_by_id.setdefault(m.group(1), []).append(m)

    n = len(top_matches)
    for i, m in enumerate(top_matches):
        ident = m.group(1)
        block_start = m.start()
        block_end = top_matches[i + 1].start() if i + 1 < n else len(body)

        ident_subs = [sm for sm in subs_by_id.get(ident, []) if block_start < sm.start() < block_end]
        next_ident = top_matches[i + 1].group(1) if i + 1 < n else None
        if not ident_subs and next_ident and re.match(rf"^{re.escape(ident)}[a-z]$", next_ident):
            continue

        question_text_end = ident_subs[0].start() if ident_subs else block_end
        raw_q = body[block_start:question_text_end]
        raw_q = re.sub(rf"^{re.escape(ident)}\.?\s+", "", raw_q, count=1)

        heuristic_split = False
        if ident_subs:
            question = " ".join(raw_q.split())
            answer = body[ident_subs[0].start():block_end].strip()
        else:
            heuristic_split = True
            collapsed = " ".join(raw_q.split())
            q_part, a_part = find_question_answer_split(collapsed)
            question = q_part.strip()
            answer = a_part.strip()
            if not answer:
                warnings.append(f"no_answer_boundary_found_for_marker_{ident}")

        question = question.strip()
        answer = answer.strip()
        if not question:
            warnings.append(f"empty_question_at_marker_{ident}")
            continue
        if not answer:
            warnings.append(f"empty_answer_for_marker_{ident}")

        pairs.append((question, answer, heuristic_split))

    return pairs, warnings


# ---------------------------------------------------------------------------
# Auto-detection
# ---------------------------------------------------------------------------

def detect_and_parse_docx(paragraphs, min_markers):
    pairs_b, warn_b, count_b = parse_marker_paragraphs(paragraphs, MARKER_RE_B)
    if count_b >= min_markers:
        return [(q, a, False) for q, a in pairs_b], warn_b, "B (docx 'Question N.')", count_b

    pairs_a, warn_a, count_a = parse_marker_paragraphs(paragraphs, MARKER_RE_A)
    if count_a >= min_markers:
        return [(q, a, False) for q, a in pairs_a], warn_a, "A (docx direct 'N.')", count_a

    msg = (f"no_qa_structure_detected: found {count_b} 'Question N.' marker(s) and "
           f"{count_a} bare 'N.' marker(s) - need >= {min_markers} of one kind to auto-detect")
    return [], [msg], None, max(count_a, count_b)


def detect_and_parse_pdf(text, min_markers):
    lines = text.split("\n")
    pairs_b, warn_b, count_b = parse_marker_paragraphs(lines, MARKER_RE_B)
    if count_b >= min_markers:
        return [(q, a, False) for q, a in pairs_b], warn_b, "B (pdf 'Question N.')", count_b

    pairs_c, warn_c = parse_family_c(text)
    count_c = len(pairs_c)
    if count_c >= min_markers:
        return pairs_c, warn_c, "C (pdf old-style numbered memo)", count_c

    msg = (f"no_qa_structure_detected: found {count_b} 'Question N.' marker(s) and "
           f"{count_c} family-C numbered question/sub-answer pair(s) - need >= {min_markers} to auto-detect")
    return [], [msg], None, max(count_b, count_c)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ingest-dir", default=DEFAULT_INGEST_DIR)
    parser.add_argument("--ingested-dir", default=DEFAULT_INGESTED_DIR)
    parser.add_argument("--xlsx", default=DEFAULT_XLSX)
    parser.add_argument("--min-markers", type=int, default=DEFAULT_MIN_MARKERS,
                         help="Minimum marker hits required before trusting an auto-detected structure")
    args = parser.parse_args()

    ingest_dir = args.ingest_dir
    ingested_dir = args.ingested_dir
    xlsx_path = args.xlsx
    min_markers = args.min_markers

    os.makedirs(ingest_dir, exist_ok=True)
    os.makedirs(ingested_dir, exist_ok=True)

    if not os.path.exists(xlsx_path):
        print(f"STOP: target workbook not found: {xlsx_path}")
        sys.exit(1)

    candidates = []
    for fname in sorted(os.listdir(ingest_dir)):
        if fname.startswith("~$") or fname.startswith("."):
            continue
        if os.path.splitext(fname)[1].lower() in (".docx", ".pdf"):
            candidates.append(fname)

    print(f"Ingest folder : {ingest_dir}")
    print(f"Ingested folder: {ingested_dir}")
    print(f"Workbook       : {xlsx_path}")
    print(f"Files found to ingest: {len(candidates)}")
    for f in candidates:
        print(f"  {f}")
    print()

    if not candidates:
        print("Nothing to ingest. No changes made.")
        return

    wb = openpyxl.load_workbook(xlsx_path)
    if SHEET_NAME not in wb.sheetnames:
        print(f"STOP: sheet '{SHEET_NAME}' not found in workbook (found {wb.sheetnames})")
        sys.exit(1)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(headers)}
    required = ["Question", "Answer", "SourceFile", "isDESA", "Category", "Year", "FilePathFull", "FolderPath"]
    missing = [h for h in required if h not in col]
    if missing:
        print(f"STOP: workbook missing expected columns: {missing} (found {headers})")
        sys.exit(1)

    last_row = ws.max_row
    existing_sources = set()
    existing_keys = set()
    for r in range(2, last_row + 1):
        q = ws.cell(row=r, column=col["Question"]).value
        a = ws.cell(row=r, column=col["Answer"]).value
        src = ws.cell(row=r, column=col["SourceFile"]).value
        if src:
            existing_sources.add(str(src).strip().lower())
        if q and a:
            existing_keys.add(norm_key(q, a))

    new_rows = []
    file_reports = []
    skipped_files = []
    moved_files = []

    for fname in candidates:
        path = os.path.join(ingest_dir, fname)

        if fname.strip().lower() in existing_sources:
            print(f"SKIP (filename already present as a SourceFile in the workbook): {fname}")
            skipped_files.append((fname, "filename already present as a SourceFile in the workbook - left in Ingest; rename if this is intentionally a revised version"))
            continue
        if os.path.exists(os.path.join(ingested_dir, fname)):
            print(f"SKIP (same filename already exists in Ingested folder): {fname}")
            skipped_files.append((fname, "a file with this exact name already exists in the Ingested folder - left in Ingest"))
            continue

        ext = os.path.splitext(fname)[1].lower()
        try:
            if ext == ".docx":
                paragraphs = extract_docx_paragraphs(path)
                pairs, warnings, family, marker_count = detect_and_parse_docx(paragraphs, min_markers)
            else:
                text, has_text = extract_pdf_text(path)
                if not has_text:
                    pairs, warnings, family, marker_count = [], ["needs_ocr: no extractable text layer"], None, 0
                else:
                    pairs, warnings, family, marker_count = detect_and_parse_pdf(text, min_markers)
        except Exception as e:
            pairs, warnings, family, marker_count = [], [f"extraction_error: {e!r}"], None, 0

        if family is None or not pairs:
            print(f"UNPARSEABLE - left in Ingest: {fname} | {'; '.join(warnings)}")
            skipped_files.append((fname, "; ".join(warnings) if warnings else "no Q&A structure auto-detected"))
            continue

        file_new = 0
        file_dup = 0
        year_val = extract_year_from_filename(fname)
        filepath_val = local_file_uri(os.path.join(ingested_dir, fname))

        for item in pairs:
            if len(item) == 3:
                q, a, heuristic = item
            else:
                q, a = item
                heuristic = False
            key = norm_key(q, a)
            if key in existing_keys:
                file_dup += 1
                continue
            existing_keys.add(key)
            new_rows.append({
                "Question": q, "Answer": a, "SourceFile": fname,
                "Category": classify_category(q, a),
                "isDESA": classify_is_desa(a),
                "Reviewed": "No", "Cleared": "",
                "Year": year_val, "FilePathFull": filepath_val,
                "FolderPath": ingested_dir, "DocumentID": "",
                "heuristic_split": heuristic,
            })
            file_new += 1

        print(f"PARSED [{family}] {fname} -> {len(pairs)} pair(s) extracted, {file_new} new, {file_dup} duplicate of existing workbook content"
              + (f" | {'; '.join(warnings)}" if warnings else ""))

        file_reports.append({
            "file": fname, "family": family, "marker_count": marker_count,
            "pairs_extracted": len(pairs), "pairs_new": file_new, "pairs_dup": file_dup,
            "warnings": warnings,
        })
        moved_files.append(fname)

    print()
    if not new_rows and not moved_files:
        print("No new Q&A pairs and no files qualified to move. No changes made to the workbook.")
        if skipped_files:
            print(f"\n{len(skipped_files)} file(s) left in Ingest:")
            for fname, reason in skipped_files:
                print(f"  {fname}: {reason}")
        return

    if new_rows:
        start_row = last_row + 1
        wrap = Alignment(wrap_text=True, vertical="top")
        flag_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

        for offset, item in enumerate(new_rows):
            r = start_row + offset
            for h, colidx in col.items():
                ws.cell(row=r, column=colidx, value=item.get(h, ""))
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).alignment = wrap
            src_cell = ws.cell(row=r, column=col["SourceFile"])
            if item["FilePathFull"]:
                src_cell.hyperlink = item["FilePathFull"]
                src_cell.font = Font(color="0563C1", underline="single")
            if item.get("heuristic_split"):
                for c in range(1, len(headers) + 1):
                    ws.cell(row=r, column=c).fill = flag_fill

        new_last_row = start_row + len(new_rows) - 1

        for tbl_name in list(ws.tables.keys()):
            del ws.tables[tbl_name]
        last_col_letter = get_column_letter(len(headers))
        tbl = Table(displayName=TABLE_NAME, ref=f"A1:{last_col_letter}{new_last_row}")
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(tbl)

        cat_col_letter = get_column_letter(col["Category"])
        desa_col_letter = get_column_letter(col["isDESA"])
        kept_dvs = []
        for dv in list(ws.data_validations.dataValidation):
            sqref_str = str(dv.sqref)
            if sqref_str.startswith(f"{cat_col_letter}2:{cat_col_letter}") or sqref_str.startswith(f"{desa_col_letter}2:{desa_col_letter}"):
                continue
            kept_dvs.append(dv)
        ws.data_validations.dataValidation = kept_dvs

        dv_cat = DataValidation(type="list", formula1="Dictionaries!$A$2:$A$10", allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv_cat)
        dv_cat.add(f"{cat_col_letter}2:{cat_col_letter}{new_last_row}")

        dv_desa = DataValidation(type="list", formula1="Dictionaries!$B$2:$B$3", allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv_desa)
        dv_desa.add(f"{desa_col_letter}2:{desa_col_letter}{new_last_row}")

    log_ws = wb["Log"] if "Log" in wb.sheetnames else wb.create_sheet("Log")
    wrap = Alignment(wrap_text=True, vertical="top")
    ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    total_new = sum(r["pairs_new"] for r in file_reports)
    total_dup = sum(r["pairs_dup"] for r in file_reports)
    log_ws.append([])
    log_ws.append([f"INGEST RUN {ts}: {total_new} new row(s) added from {len(moved_files)} file(s); "
                   f"{total_dup} pair(s) skipped as duplicates of existing workbook content; "
                   f"{len(skipped_files)} file(s) left in Ingest (see below)."])
    for r in file_reports:
        warn_str = "; ".join(r["warnings"]) if r["warnings"] else "-"
        log_ws.append(["", r["file"], r["family"], f"{r['pairs_new']} new / {r['pairs_dup']} dup of {r['pairs_extracted']} extracted", warn_str])
    if skipped_files:
        log_ws.append(["", "Files left in Ingest (not moved):"])
        for fname, reason in skipped_files:
            log_ws.append(["", fname, reason])
    if new_rows:
        log_ws.append(["", "NOTE: Year parsed from filename (first 19xx/20xx run found) or left blank if none. "
                           "FilePathFull/FolderPath point to the local Ingested-folder location (file:// link) - "
                           "no SharePoint URL or DocumentID exists yet for newly ingested files, so DocumentID is "
                           "left blank rather than invented. Category/isDESA assigned by the same automated "
                           "keyword heuristic used elsewhere in this workbook - spot-check before relying on them."])
    for row in log_ws.iter_rows(min_row=log_ws.max_row - len(file_reports) - len(skipped_files) - 3, max_row=log_ws.max_row):
        for cell in row:
            cell.alignment = wrap

    wb.save(xlsx_path)

    for fname in moved_files:
        shutil.move(os.path.join(ingest_dir, fname), os.path.join(ingested_dir, fname))

    print("=" * 100)
    print("SUMMARY")
    print("=" * 100)
    for r in file_reports:
        print(f"  {r['file']:<60} [{r['family']}] {r['pairs_new']} new / {r['pairs_dup']} dup of {r['pairs_extracted']} extracted")
    print(f"\nTotal new rows added to workbook: {total_new}")
    print(f"Total duplicate pairs skipped: {total_dup}")
    print(f"Files moved to Ingested: {len(moved_files)}")
    if skipped_files:
        print(f"\nFiles left in Ingest ({len(skipped_files)}):")
        for fname, reason in skipped_files:
            print(f"  {fname}: {reason}")
    else:
        print("\nNo files left in Ingest.")


if __name__ == "__main__":
    main()
