"""
create_sharepoint_sample.py

Generates a small sample Excel file (one row per record type) for creating
the SharePoint list via "Import from Excel". SharePoint's import is capped
at 5,000 rows, so the full country_data_sharepoint.xlsx cannot be used for
initial list creation when the dataset is large.

Workflow:
  1. Run parse_country_data.py  -> country_data_sharepoint.xlsx  (full data)
  2. Run this script            -> country_data_sharepoint_sample.xlsx  (6 rows)
  3. Import the sample into SharePoint to create the list and its columns
  4. Fix column types in SharePoint List Settings (see README.md)
  5. Run the Power Automate flow to load all rows from the full Excel

Usage:
    python create_sharepoint_sample.py
"""

from pathlib import Path
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

INPUT_PATH  = Path(__file__).parent / "country_data_sharepoint.xlsx"
OUTPUT_PATH = Path(__file__).parent / "country_data_sharepoint_sample.xlsx"


def main():
    df = pd.read_excel(INPUT_PATH, dtype=str)
    print(f"Full dataset: {len(df)} rows, {len(df.columns)} columns")

    # One row per record type — enough for SharePoint to see all columns
    sample = df.groupby("record_type").first().reset_index()
    print(f"Sample: {len(sample)} rows (one per record type)")
    print(sample[["record_type", "country_code"]].to_string(index=False))

    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        sample.to_excel(writer, sheet_name="All Records", index=False)
        ws = writer.sheets["All Records"]
        last_col = get_column_letter(ws.max_column)
        ref = f"A1:{last_col}{ws.max_row}"
        tbl = Table(displayName="AllRecords", ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(tbl)

    print(f"\nSample written -> {OUTPUT_PATH}")
    print("Import this file into SharePoint to create the list, then run the Power Automate flow.")


if __name__ == "__main__":
    main()
