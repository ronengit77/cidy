"""
parse_country_data.py

Parses Cidy Country Data .docx files into a flat SharePoint-ready Excel table.
Run this script whenever source .docx files are updated, then trigger the
Power Automate "Country Data Refresh" flow to reload the SharePoint list.

Requirements:
    pip install python-docx openpyxl pandas

Usage:
    python parse_country_data.py                  # process all files
    python parse_country_data.py jam-prjct-...docx  # test a single file
"""

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ── CONFIG ────────────────────────────────────────────────────────────────────

DOCX_FOLDER = Path(__file__).parent.parent / "Knowledge" / "Programme Development" / "Cidy - Country Data"
OUTPUT_PATH = Path(__file__).parent / "country_data_sharepoint.xlsx"

# ── HELPERS ───────────────────────────────────────────────────────────────────

def _text(element):
    return "".join(n.text or "" for n in element.iter() if n.tag == qn("w:t")).strip()


def _table_rows(tbl_element, doc):
    from docx.table import Table
    table = Table(tbl_element, doc)
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def _parse_date(val):
    if not val:
        return None
    val = str(val).strip()
    if val.lower() in ("na", "n/a", "tbd", "ongoing", ""):
        return None
    val = re.sub(r"\bSept\b", "Sep", val, flags=re.IGNORECASE)
    for fmt in ("%d %b %Y", "%d %B %Y", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(val, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _clean_money(val):
    if not val:
        return None
    val = re.sub(r"[^\d.]", "", str(val))
    try:
        return float(val)
    except ValueError:
        return None


def _clean_int(val):
    if val is None:
        return None
    try:
        return int(re.sub(r"[^\d]", "", str(val)))
    except ValueError:
        return None


SECTIONS = {
    "key metrics":     "key_metrics",
    "projects":        "projects",
    "requests":        "requests",
    "rptc proposals":  "proposals",
    "rptc activities": "activities",
    "unct":            "contacts",
    "contacts":        "contacts",
    "data quality":    "data_quality",
}


def _parse_metrics(rows):
    metrics = {}
    for row in rows:
        if len(row) >= 2 and row[0] and row[0].lower() != "metric":
            metrics[row[0].strip()] = row[1].strip()
    return metrics


def _detect_section(text):
    lower = text.lower().strip()
    for key, name in SECTIONS.items():
        if lower.startswith(key):
            return name
    return None


def _header_map(header_row):
    return {cell.lower().strip(): i for i, cell in enumerate(header_row)}


def _get(row, col_map, *keys):
    for key in keys:
        for col, idx in col_map.items():
            if key in col and idx < len(row):
                return row[idx].strip()
    return None


# ── SECTION PARSERS ───────────────────────────────────────────────────────────

def _parse_projects(rows):
    if len(rows) < 2:
        return []
    cols = _header_map(rows[0])
    results = []
    for row in rows[1:]:
        if not any(c.strip() for c in row):
            continue
        results.append({
            "title":             _get(row, cols, "title"),
            "division":          _get(row, cols, "division"),
            "status":            _get(row, cols, "status"),
            "start_date":        _parse_date(_get(row, cols, "start")),
            "end_date":          _parse_date(_get(row, cols, "end")),
            "budget":            _clean_money(_get(row, cols, "budget")),
            "partner_countries": _get(row, cols, "countr"),
        })
    return results


def _parse_requests(rows):
    if len(rows) < 2:
        return []
    cols = _header_map(rows[0])
    results = []
    for row in rows[1:]:
        if not any(c.strip() for c in row):
            continue
        results.append({
            "ref_number":   (_get(row, cols, "ref") or "").replace(",", "") or None,
            "title":        _get(row, cols, "title"),
            "status":       _get(row, cols, "status"),
            "division":     _get(row, cols, "division"),
            "request_date": _parse_date(_get(row, cols, "request")),
            "start_date":   _parse_date(_get(row, cols, "start")),
            "end_date":     _parse_date(_get(row, cols, "end")),
            "contact":      _get(row, cols, "contact"),
            "summary":      _get(row, cols, "summary", "description"),
        })
    return results


def _parse_proposals(rows):
    if len(rows) < 2:
        return []
    cols = _header_map(rows[0])
    results = []
    for row in rows[1:]:
        if not any(c.strip() for c in row):
            continue
        results.append({
            "project_id":        _get(row, cols, "project id", "id"),
            "title":             _get(row, cols, "title"),
            "division":          _get(row, cols, "division"),
            "planned_start":     _parse_date(_get(row, cols, "start")),
            "planned_end":       _parse_date(_get(row, cols, "end")),
            "budget":            _clean_money(_get(row, cols, "budget")),
            "partner_countries": _get(row, cols, "countr"),
        })
    return results


def _parse_activities(rows):
    if len(rows) < 2:
        return []
    cols = _header_map(rows[0])
    results = []
    for row in rows[1:]:
        if not any(c.strip() for c in row):
            continue
        results.append({
            "project_id":         _get(row, cols, "project id", "id"),
            "title":              _get(row, cols, "title"),
            "division":           _get(row, cols, "division"),
            "start_date":         _parse_date(_get(row, cols, "start")),
            "end_date":           _parse_date(_get(row, cols, "end")),
            "consumed_budget":    _clean_money(_get(row, cols, "consumed", "budget")),
            "participants":       _clean_int(_get(row, cols, "participant")),
            "women_participants": _clean_int(_get(row, cols, "women")),
        })
    return results


# ── MAIN PARSER ───────────────────────────────────────────────────────────────

def parse_docx(filepath):
    doc = Document(filepath)
    filename = Path(filepath).name
    country_code = filename.split("-prjct-")[0].upper()

    result = {
        "country_code":  country_code,
        "filename":      filename,
        "country_name":  None,
        "region":        None,
        "country_type":  None,
        "snapshot_date": None,
        "export_date":   None,
        "metrics":       {},
        "projects":      [],
        "requests":      [],
        "proposals":     [],
        "activities":    [],
    }

    body_items = []
    for child in doc.element.body:
        tag = child.tag.split("}")[-1]
        if tag == "p":
            body_items.append(("para", _text(child)))
        elif tag == "tbl":
            body_items.append(("table", _table_rows(child, doc)))

    # Parse header: country name, region, country type, dates
    header_paras = []
    for kind, content in body_items[:15]:
        if kind != "para" or not content:
            continue
        if "Exported" in content and "Snapshot" in content:
            m = re.search(r"Exported\s+(.+?)\s*\|", content)
            if m:
                result["export_date"] = m.group(1).strip()
            m = re.search(r"Snapshot\s+(.+?)(?:\s*\(|$)", content)
            if m:
                result["snapshot_date"] = m.group(1).strip()
            type_keywords   = {"SIDS", "LDC", "LLDC"}
            region_keywords = {"Americas", "Africa", "Asia", "Europe", "Oceania"}
            if len(header_paras) >= 2:
                region_line = header_paras[-1]
                name_line   = " ".join(header_paras[:-1])
            else:
                region_line = ""
                name_line   = " ".join(header_paras)
            rtokens = [t for t in region_line.split() if t != "|"]
            result["country_type"] = " ".join(t for t in rtokens if t in type_keywords) or None
            result["region"]       = " ".join(t for t in rtokens if t in region_keywords) or None
            ntokens = [t for t in name_line.split()
                       if t not in type_keywords and t not in region_keywords and t != "|"]
            result["country_name"] = " ".join(ntokens).strip() or country_code
            break
        header_paras.append(content)

    # Walk body in document order, parse each section table
    current_section = None
    for kind, content in body_items:
        if kind == "para" and content:
            section = _detect_section(content)
            if section:
                current_section = section
        elif kind == "table":
            if not content or current_section is None:
                continue
            if current_section == "key_metrics":
                result["metrics"]    = _parse_metrics(content)
            elif current_section == "projects":
                result["projects"]   = _parse_projects(content)
            elif current_section == "requests":
                result["requests"]   = _parse_requests(content)
            elif current_section == "proposals":
                result["proposals"]  = _parse_proposals(content)
            elif current_section == "activities":
                result["activities"] = _parse_activities(content)

    return result


# ── FLAT LIST BUILDER ─────────────────────────────────────────────────────────

def build_flat_list(all_results):
    """One row per record (Project / Request / Proposal / Activity). Country fields repeated on every row."""
    rows = []

    for r in all_results:
        ctx = {
            "country_code":  r["country_code"],
            "country_name":  r["country_name"],
            "region":        r["region"],
            "country_type":  r["country_type"],
            "snapshot_date": r["snapshot_date"],
        }

        for p in r["projects"]:
            rows.append({**ctx, "record_type": "Project",
                "title": p.get("title"), "division": p.get("division"),
                "status": p.get("status"), "start_date": p.get("start_date"),
                "end_date": p.get("end_date"), "budget": p.get("budget"),
                "partner_countries": p.get("partner_countries"),
                "ref_number": None, "request_date": None, "contact": None, "summary": None,
                "project_id": None, "consumed_budget": None,
                "participants": None, "women_participants": None,
            })

        for req in r["requests"]:
            rows.append({**ctx, "record_type": "Request",
                "title": req.get("title"), "division": req.get("division"),
                "status": req.get("status"), "start_date": req.get("start_date"),
                "end_date": req.get("end_date"), "budget": None,
                "partner_countries": None,
                "ref_number": req.get("ref_number"), "request_date": req.get("request_date"),
                "contact": req.get("contact"), "summary": req.get("summary"),
                "project_id": None, "consumed_budget": None,
                "participants": None, "women_participants": None,
            })

        for prop in r["proposals"]:
            rows.append({**ctx, "record_type": "Proposal",
                "title": prop.get("title"), "division": prop.get("division"),
                "status": None, "start_date": prop.get("planned_start"),
                "end_date": prop.get("planned_end"), "budget": prop.get("budget"),
                "partner_countries": prop.get("partner_countries"),
                "ref_number": None, "request_date": None, "contact": None, "summary": None,
                "project_id": prop.get("project_id"), "consumed_budget": None,
                "participants": None, "women_participants": None,
            })

        for act in r["activities"]:
            rows.append({**ctx, "record_type": "Activity",
                "title": act.get("title"), "division": act.get("division"),
                "status": None, "start_date": act.get("start_date"),
                "end_date": act.get("end_date"), "budget": None,
                "partner_countries": None,
                "ref_number": None, "request_date": None, "contact": None, "summary": None,
                "project_id": act.get("project_id"),
                "consumed_budget": act.get("consumed_budget"),
                "participants": act.get("participants"),
                "women_participants": act.get("women_participants"),
            })

    df = pd.DataFrame(rows)
    for col in ("ref_number", "project_id"):
        df[col] = df[col].where(df[col].isna(), df[col].astype(str))
    return df


# ── EXCEL OUTPUT ──────────────────────────────────────────────────────────────

def write_excel(df_flat, output_path):
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_flat.to_excel(writer, sheet_name="All Records", index=False)
        ws = writer.sheets["All Records"]
        last_col = get_column_letter(ws.max_column)
        ref = f"A1:{last_col}{ws.max_row}"
        tbl = Table(displayName="AllRecords", ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(tbl)
    print(f"SharePoint Excel written -> {output_path}")


# ── ENTRY POINT ───────────────────────────────────────────────────────────────

def main(single_file=None):
    if single_file:
        docx_files = [Path(single_file)]
    else:
        docx_files = sorted(DOCX_FOLDER.glob("*.docx"))
    print(f"Parsing {len(docx_files)} file(s)\n")

    all_results, errors = [], []
    for f in docx_files:
        try:
            r = parse_docx(f)
            all_results.append(r)
            print(f"  {r['country_code']:12s} projects={len(r['projects']):3d}  "
                  f"requests={len(r['requests']):3d}  proposals={len(r['proposals']):3d}  "
                  f"activities={len(r['activities']):3d}")
        except Exception as e:
            errors.append((f.name, str(e)))
            print(f"  ERROR {f.name}: {e}")

    if errors:
        print(f"\n{len(errors)} file(s) failed:")
        for name, msg in errors:
            print(f"  {name}: {msg}")

    df_flat = build_flat_list(all_results)
    write_excel(df_flat, OUTPUT_PATH)

    print(f"\nTotal rows: {len(df_flat)}")
    print(df_flat["record_type"].value_counts().to_string())


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        main(single_file=DOCX_FOLDER / sys.argv[1])
    else:
        main()
